#!/usr/bin/env python3
"""Referto del progetto in risultati/referto/: referto.html, referto.xlsx, referto.pdf (contratto v0.2, D036).

Regole: italiano semplice con i termini tecnici spiegati (D027); MetaPhlAn e Kraken SEMPRE in tabelle e fogli
separati, mai fusi; Kraken mostra solo i microbi (batteri, archei, funghi), con Homo, Leishmania e il resto
segnalati a parte (D023); parametri usati con la fonte (D022); riquadro "Strumenti e database usati" (D036);
virus come opzione sperimentale con postilla (D017, D033); sicurezza alimentare se calcolata (risultati/sicurezza/).

Gira con il Python di pmn_reports_py: libreria standard + openpyxl (Excel) + PySide6/QtGui (PDF, senza schermo:
QTextDocument + QPdfWriter, dallo stesso HTML). Se manca un pezzo, gli altri file si fanno lo stesso e
l'uscita lo dice (il motore lo trasforma nell'avviso REFERTO_PARZIALE).

Uso: referto.py --risultati DIR [--progetto config.yaml]   (lingua dei testi: PMN_LINGUA, motore/lingua.py)
stdout: una riga JSON {"html", "xlsx", "pdf", "problemi": [...]}. Codici: 0 fatto (anche parziale), 1 risultati mancanti.
"""
import argparse
import csv
import html
import json
import os
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "motore"))
import diversita  # noqa: E402
import lingua  # noqa: E402
import multicampione  # noqa: E402
import system_one  # noqa: E402
from lingua import N_, _  # noqa: E402

OSPITI = {"sus_scrofa": N_("maiale (Sscrofa11.1)"), "bos_taurus": N_("bovino (ARS-UCD1.2)"),
          "ovis_aries": N_("pecora (ARS-UI_Ramb_v2.0)"), "grch38": N_("uomo (GRCh38)"),
          "homo_sapiens": N_("uomo (GRCh38)")}
MICROBI = ("batteri", "archei", "funghi")   # categorie: chiavi, si traducono solo quando si mostrano
CATEGORIE = [N_("batteri"), N_("archei"), N_("funghi"), N_("virus"), N_("altri eucarioti"), N_("altro"),
             N_("buono"), N_("da controllare"), N_("da rifare"), N_("migliorata"), N_("fssp")]
DOMINANTI = 15   # specie in sintesi.*.dominanti (v0.3); la frase ne nomina 3
VERSIONI = {"metaphlan": "MetaPhlAn", "database_metaphlan": N_("Database di MetaPhlAn"), "kraken2": "Kraken 2",
            "bracken": "Bracken", "database_kraken": N_("Database di Kraken"), "bowtie2": N_("Bowtie 2 (ospite, fette)"),
            "fastp": N_("fastp (pulizia delle read)"), "samtools": "SAMtools",
            "fastqc": N_("FastQC (qualità delle read)"), "multiqc": N_("MultiQC (riepilogo del controllo qualità)")}
MODELLI = {"listeria": N_("modello della Listeria"), "listeria_lab": N_("modello della Listeria con i batteri lattici"),
           "lab": N_("modello dei batteri lattici")}
NOMI_PRODOTTO = {"temperatura": N_("temperatura {} °C"), "sale_fase_acquosa": N_("sale nella fase acquosa {} %"),
                 "aw": "aw {}", "ph": "pH {}", "fenoli": N_("fenoli del fumo {} ppm"), "co2": "CO2 {} %",
                 "nitrito": N_("nitrito {} ppm")}
POSTILLA_VIRUS = N_("Virus: opzione sperimentale, non validata con il database diviso in 6 parti. Nei pareggi fra copie "
                    "virali quasi identiche la scelta può cambiare rispetto al MetaPhlAn standard (D017, D033).")
GLOSSARIO = [
    (N_("Read (lettura)"), N_("un pezzetto di DNA letto dal sequenziatore. Con le coppie (read pairs) ogni frammento è "
                              "letto dai due lati.")),
    (N_("Ospite"), N_("l'animale da cui viene il prodotto (es. maiale) e l'uomo: il loro DNA si toglie prima di cercare "
                      "i taxa.")),
    (N_("Abbondanza relativa"), N_("la percentuale di DNA di una specie rispetto al totale: non è il numero di cellule "
                                   "vive per grammo.")),
    ("MetaPhlAn", N_("riconosce le specie da geni caratteristici di ognuna (marcatori): preciso sulle specie note.")),
    ("Kraken + Bracken", N_("confronta ogni read con interi genomi di riferimento, poi stima le abbondanze: vede più "
                            "specie, anche funghi, ma con più rumore. Per questo i due risultati stanno in tabelle "
                            "separate.")),
    (N_("Soglia di abbondanza (solo Kraken)"), N_("sotto questa percentuale un taxon di Kraken non si mostra nelle "
                                                  "tabelle del referto (resta nell'Excel). MetaPhlAn non ha soglia: si "
                                                  "mostrano tutte le specie.")),
    (N_("Diversità alfa"), N_("quanti taxa diversi ci sono in un campione (ricchezza) e quanto le loro abbondanze sono "
                              "simili fra loro (Shannon, Simpson).")),
    (N_("Rarefazione"), N_("confrontare i campioni alla stessa profondità, cioè con lo stesso numero di read, "
                           "togliendo a caso le read in più dai campioni più ricchi.")),
]
STATO_DIVERSITA = {"intero": N_("tutte le read"), "pareggiato": N_("profondità uguale"),
                   "non_pareggiato": N_("non pareggiato")}


# ---------------- lettura dei risultati ----------------
def nome_specie(clade):
    return clade.split("|")[-1][3:].replace("_", " ")


def leggi_metaphlan(f):
    """([(specie, %)], % UNCLASSIFIED): righe s__ senza t__, percentuali come le scrive MetaPhlAn."""
    specie, non_class = [], 0.0
    for riga in Path(f).read_text(encoding="utf-8").splitlines():
        c = riga.split("\t")
        if riga.startswith("#") or len(c) < 3:
            continue
        if c[0] == "UNCLASSIFIED":
            non_class = float(c[2])
        elif c[0].split("|")[-1].startswith("s__"):
            specie.append((nome_specie(c[0]), float(c[2])))
    return sorted(specie, key=lambda x: -x[1]), non_class


def categorie_kraken(k2report):
    """{taxid: categoria} dall'albero del report di Kraken (rientro di 2 spazi per livello)."""
    out, pila = {}, []
    for riga in Path(k2report).read_text(encoding="utf-8").splitlines():
        c = riga.split("\t")
        if len(c) < 6:
            continue
        nome, taxid = c[-1], c[-2].strip()
        livello = (len(nome) - len(nome.lstrip(" "))) // 2
        while pila and pila[-1][0] >= livello:
            pila.pop()
        pila.append((livello, taxid))
        antenati = {t for _, t in pila}
        if "10239" in antenati:
            out[taxid] = "virus"
        elif "2" in antenati:
            out[taxid] = "batteri"
        elif "2157" in antenati:
            out[taxid] = "archei"
        elif "2759" in antenati:
            out[taxid] = "funghi" if "4751" in antenati else "altri eucarioti"
        else:
            out[taxid] = "altro"
    return out


def leggi_kraken(bracken, k2report):
    """Specie di Bracken con categoria, % sul totale assegnato e % fra i soli microbi."""
    cat = categorie_kraken(k2report) if Path(k2report).is_file() else {}
    righe = []
    with open(bracken, encoding="utf-8", newline="") as f:
        for r in csv.DictReader(f, delimiter="\t"):
            righe.append({"nome": r["name"], "taxid": r["taxonomy_id"], "categoria": cat.get(r["taxonomy_id"], "altro"),
                          "letture": int(r["new_est_reads"])})
    # % dalle read intere, come il passo 33 (beta): fraction_total_reads ha 5 decimali (passi dello 0,001 %), così un
    # taxon allo 0,00455 % diventava 0,005 % e passava la soglia rapido (P29)
    totale = sum(r["letture"] for r in righe)
    for r in righe:
        r["pct"] = 100.0 * r["letture"] / totale if totale else 0.0
    somma = sum(r["letture"] for r in righe if r["categoria"] in MICROBI)   # intera: stesso numero su ogni Python
    for r in righe:
        r["pct_microbi"] = 100.0 * r["letture"] / somma if r["categoria"] in MICROBI and somma else None
    return sorted(righe, key=lambda r: -r["pct"])


def tsv_dict(f):
    out = {}
    if Path(f).is_file():
        for riga in Path(f).read_text(encoding="utf-8").splitlines()[1:]:
            k, _, v = riga.partition("\t")
            out[k] = v
    return out


def profili_kraken():
    out = {}
    for riga in (ROOT / "refs" / "kraken_profili.tsv").read_text(encoding="utf-8").splitlines():
        if riga.strip() and not riga.startswith("#"):
            c = riga.split("\t")
            out[c[0]] = {"confidence": c[1], "min_reads": c[2], "soglia": c[3], "fonte": c[4]}
    return out


def numero(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def leggi_campione(risultati, file):
    """{"metaphlan", "non_classificato", "kraken"} dai file di un campione (None se il file manca)."""
    d = {"metaphlan": None, "non_classificato": None, "kraken": None}
    if file.get("metaphlan") and (risultati / file["metaphlan"]).is_file():
        d["metaphlan"], d["non_classificato"] = leggi_metaphlan(risultati / file["metaphlan"])
    if file.get("kraken") and (risultati / file["kraken"]).is_file():
        b = risultati / file["kraken"]
        d["kraken"] = leggi_kraken(b, b.with_suffix(".k2report"))
    return d


def elenca(nomi_pct, base):
    nomi_pct = list(nomi_pct)
    testo = [f"{n} ({pct(v, 1)} %{base})" for n, v in nomi_pct]
    return _(" e ").join(testo) if len(testo) <= 2 else ", ".join(testo[:-1]) + _(" e ") + testo[-1]


def articolo(numero_scritto):
    """Articolo italiano davanti a una percentuale, come si legge: "Lo 0,0", "L'8,1", "L'11", "Il 53,7"."""
    if lingua.attuale() != "it":
        return ""
    intero = numero_scritto.split(",")[0].split(".")[0]
    if intero == "0":
        return "Lo "
    return "L'" if intero in ("1", "11") or intero.startswith("8") else "Il "


def sintesi(d, soglia):
    """Sintesi di un campione per risultati.json e per "In breve" (l'interfaccia mostra, non calcola: D006).
    MetaPhlAn e Kraken restano separati. MetaPhlAn: tutte le specie, nessuna soglia; Kraken: solo sopra la soglia."""
    out = {}
    if d.get("metaphlan") is not None:
        specie = d["metaphlan"]
        nc = d.get("non_classificato") or 0.0
        frase = (_("MetaPhlAn: {n} specie trovate", n=len(specie)) +
                 (_("; dominano {specie}.", specie=elenca(specie[:3], _(" delle read"))) if specie else ".") +
                 _(" {art}{pct} % delle read non è assegnato a specie note.", art=articolo(pct(nc, 1)), pct=pct(nc, 1)))
        out["metaphlan"] = {"specie": len(specie), "dominanti": [{"nome": n, "pct": v} for n, v in specie[:DOMINANTI]],
                            "non_assegnata_pct": nc, "frase": frase}
    if d.get("kraken") is not None:
        micro = [r for r in d["kraken"] if r["categoria"] in MICROBI and r["pct"] >= soglia]
        altri = [r for r in d["kraken"] if r["categoria"] not in MICROBI]
        frase = (_("Kraken: {n} taxa sopra la soglia", n=len(micro)) +
                 (_("; dominano {specie}.", specie=elenca([(r['nome'], r['pct_microbi']) for r in micro[:3]],
                                                           _(" fra i taxa"))) if micro else ".") +
                 (_(" Segnalati a parte (non del prodotto, o virus): ") +
                  ", ".join(f"{r['nome']} {pct(r['pct'], 3)} %" for r in altri) + "." if altri else ""))
        out["kraken"] = {"microbi": len(micro),
                         "dominanti": [{"nome": r["nome"], "pct_microbi": r["pct_microbi"]} for r in micro[:DOMINANTI]],
                         "non_microbi": [{"nome": r["nome"], "gruppo": r["categoria"], "pct": r["pct"],
                                          "letture": r["letture"]} for r in altri],
                         "frase": frase}
    return out


def raccogli(risultati, cfg):
    """Tutto quello che va nel referto, già letto dai file (una sola fonte per HTML, Excel e PDF)."""
    indice = json.loads((risultati / "risultati.json").read_text(encoding="utf-8"))
    par = tsv_dict(risultati / "tabelle" / "parameters.tsv") or {k: str(v) for k, v in indice.get("parametri", {}).items()}
    soglia = numero(par.get("abundance_threshold_pct")) or 0.0
    campioni, mpa, krk = [], {}, {}
    for c in indice.get("campioni", []):
        file = c.get("file", {})
        d = {"nome": c["nome"], "coppie": c.get("coppie_dopo_ospite"), **leggi_campione(risultati, file)}
        d["sintesi"] = sintesi(d, soglia)
        campioni.append(d)
        if d["metaphlan"] is not None:   # diversità alfa (R1): MetaPhlAn e Kraken separati, mai uniti
            mpa[d["nome"]] = (d["metaphlan"], diversita.letture_profilo(risultati / file["metaphlan"]))
        if d["kraken"] is not None:
            krk[d["nome"]] = d["kraken"]
    pari = {f.name[:-len("_profile.txt")]: (leggi_metaphlan(f)[0], diversita.letture_profilo(f))
            for f in sorted((risultati / "metaphlan_pari").glob("*_profile.txt"))}
    triage, _ = system_one.triage(risultati / "tabelle" / "triage.tsv", system_one.soglie_triage(cfg))
    sicurezza = [json.loads(f.read_text(encoding="utf-8")) for f in sorted((risultati / "sicurezza").glob("*.json"))]
    motore = cfg.get("motore") if isinstance(cfg.get("motore"), dict) else {}
    return {"progetto": str(cfg.get("project_name") or risultati.parent.name), "corsa": indice.get("corsa", ""),
            "motore": indice.get("motore", ""), "campioni": campioni, "parametri": par, "soglia": soglia,
            "versioni": indice.get("versioni", {}), "citazioni": list(indice.get("citazioni", [])),
            "triage": triage, "sicurezza": sicurezza, "virus": bool(motore.get("virus", False)),
            "qc": indice.get("qc") or {},
            "diversita": {"metaphlan": diversita.tabella_metaphlan(mpa, pari),
                          "kraken": diversita.tabella_kraken(krk, soglia, MICROBI)}}


def valori_diversita(r):
    """[read usate, ricchezza, Shannon, Simpson] di una riga di diversita.tabella_*, già scritti per le persone: "183.241
    su 2.731.766" se il campione è stato pareggiato; la ricchezza attesa della rarefazione può avere decimali."""
    ric, usate = r["ricchezza"], lingua.migliaia(f"{r['letture_usate']:,}")
    if r["letture_usate"] != r["letture"]:
        usate = _("{usate} su {tutte}", usate=usate, tutte=lingua.migliaia(f"{r['letture']:,}"))
    return [usate, str(int(ric)) if float(ric).is_integer() else pct(ric, 1), pct(r["shannon"], 2), pct(r["simpson"], 3)]


def html_diversita(d):
    """Sezione "Diversità alfa" (R1): avviso sulla profondità, poi MetaPhlAn e Kraken in due tabelle separate."""
    div = d["diversita"]
    multi = (d.get("multicampione") or {}).get("attiva")
    if not div["metaphlan"] and not div["kraken"]:
        return [f"<h2>4. {_('Diversità e confronto fra i campioni')}</h2>"] + html_multicampione(d) if multi else []
    uno = len(d["campioni"]) == 1
    titolo = _("Diversità e confronto fra i campioni") if multi else _("Diversità alfa (quanti taxa e quanto sono "
                                                                         "distribuiti)")
    h = [f"<h2>4. {titolo}</h2>",
         riquadro(_("<b>Attenzione alla profondità.</b> La ricchezza (quanti taxa si trovano) cresce con il numero di "
                    "read: con più read si trovano più taxa rari, anche nello stesso prodotto.") + " " +
                  (_("Questo progetto ha un solo campione: i valori sono alla sua profondità. Confrontali solo con "
                     "campioni analizzati alla stessa profondità.") if uno else
                   _("Per confrontare i campioni, qui i valori sono calcolati alla stessa profondità, quella del "
                     "campione con meno read (colonne \"Read usate\" e \"Letture usate\"). Confrontali solo con "
                     "campioni analizzati alla stessa profondità."))),
         '<p class="piccolo">' + _("<b>Ricchezza</b> (N. specie, N. taxa): quanti taxa diversi. <b>Shannon</b>: cresce "
                                   "con il numero di taxa e con quanto sono distribuiti in modo uniforme (0 = un solo "
                                   "taxon). "
                                   "<b>Simpson</b> (da 0 a 1): la probabilità che due read prese a caso vengano da "
                                   "taxa diversi.") + "</p>"]
    if div["metaphlan"]:
        h.append("<h3>MetaPhlAn</h3>")
        h.append(tabella([_("Campione"), _("Read usate"), _("N. specie"), "Shannon", "Simpson"],
                         [[e(r["campione"])] + valori_diversita(r) for r in div["metaphlan"]], (1, 2, 3, 4),
                         [24, 26, 16, 17, 17]))
        h.append('<p class="piccolo">' + _("Read entrate in MetaPhlAn, dopo aver tolto l'ospite. Nessuna soglia di "
                                           "abbondanza. Per i campioni con più read: un sottocampione casuale delle "
                                           "read (seme 1992), poi MetaPhlAn rifatto su quelle (una sola estrazione).")
                 + "</p>")
        non = [r["campione"] for r in div["metaphlan"] if r["stato"] == "non_pareggiato"]
        if non:
            h.append('<p class="piccolo">' + _("<b>Non pareggiato</b> ({nomi}): mancano i file per rifare MetaPhlAn su "
                                               "meno read (analisi fatta prima del 18/10/2026). I valori sono alla "
                                               "profondità del campione: non confrontarli con gli altri.",
                                               nomi=e(", ".join(non))) + "</p>")
    if div["kraken"]:
        h.append("<h3>Kraken + Bracken</h3>")
        h.append('<p class="piccolo">' + _("<b>Prototipo, non per decisioni</b> (come la sezione 3).") + "</p>")
        h.append(tabella([_("Campione"), _("Letture usate"), _("N. taxa"), "Shannon", "Simpson"],
                         [[e(r["campione"])] + valori_diversita(r) for r in div["kraken"]], (1, 2, 3, 4),
                         [24, 26, 16, 17, 17]))
        h.append('<p class="piccolo">' + _("Solo batteri, archei e funghi sopra la soglia di Kraken ({soglia} %), con "
                                           "le letture stimate da Bracken. A profondità uguale si usa il valore atteso "
                                           "della rarefazione (la media di tutti i sottocampioni possibili): per questo "
                                           "la ricchezza può avere decimali.", soglia=pct(d["soglia"], 3)) + "</p>")
    h.append('<p class="piccolo">' + _("Metodo: rarefazione a profondità uguale (Schloss 2024, mSphere "
                                       "9:e00354-23). La ricchezza osservata sottostima quella vera (Willis 2019, "
                                       "Front Microbiol 10:2407).") + "</p>")
    return h + (html_multicampione(d) if multi else [])


INDICI = {"richness": N_("Ricchezza"), "shannon": "Shannon", "simpson": "Simpson"}
DISTANZE = {"aitchison": N_("Aitchison (CLR)"), "braycurtis": "Bray-Curtis"}
LIMITI_MULTI = N_("<b>Limiti.</b> Un fattore alla volta: niente modelli con più fattori, interazioni o covariate (le "
                  "colonne in più di metadati.tsv si riportano ma non si usano). Con meno di 3 unità per gruppo non si "
                  "fa nessun test. Con pochi campioni un risultato non significativo non vuol dire che i gruppi siano "
                  "uguali. La PERMANOVA risente delle differenze di dispersione: per questo ha sempre accanto PERMDISP. "
                  "I valori p non sono corretti fra indici, distanze e confronti diversi. Le abbondanze sono DNA "
                  "relativo, non cellule vive per grammo.")


NOTE_R = {"all values are equal: no test possible": N_("valori tutti uguali: nessun test possibile")}


def nota_r(nota):
    """Una nota di multicampione.R (in inglese, come le tabelle) nella lingua del referto."""
    return _(NOTE_R[nota]) if nota in NOTE_R else nota


def citazioni_multi(d):
    """Le fonti dei metodi dell'analisi multicampione, solo se è stata fatta (almeno 2 campioni)."""
    m = d.get("multicampione") or {}
    return list(m.get("citazioni", [])) if m.get("attiva") else []


def righe_multi_excel(p):
    """[righe] del foglio di un profiler: test alfa e beta come li ha scritti multicampione.R."""
    out = [[r["comparison"], _("alfa"), _(INDICI.get(r["metric"], r["metric"])), _(multicampione.TEST[r["test"]]),
            r["n"], r["statistic_name"], r["statistic"], r["p"], _(multicampione.EFFETTO.get(r["effect_size_name"], "")),
            r["effect_size"], "", nota_r(r["note"])] for r in p["alfa"]]
    out += [[r["comparison"], _("beta"), _(DISTANZE.get(r["distance"], r["distance"])), r["test"], r["n"], "F", r["F"],
             r["p"], "R²" if r["R2"] else "", r["R2"], r["permutations"],
             nota_dispersione(r) if r["test"] == "PERMANOVA" else r["note"]] for r in p["beta"]]
    return out


def p_testo(v, cifre=None):
    """Un numero scritto da multicampione.R, nella lingua scelta; vuoto se manca. Senza cifre è un valore p: 3
    decimali, "< 0,001" sotto (stretto anche nelle colonne del PDF)."""
    x = numero(v)
    if x is None:
        return ""
    if cifre is None:
        return "&lt; " + pct(0.001, 3) if x < 0.001 else pct(x, 3)
    return pct(x, cifre)


def nota_dispersione(r):
    if r.get("dispersion_differs") == "yes":
        return _("dispersione diversa fra i gruppi: la PERMANOVA può dipendere da questa, non solo dalla composizione")
    return _("dispersione simile fra i gruppi") if r.get("dispersion_differs") == "no" else e(r.get("note", ""))


def html_multicampione(d):
    """Sottosezione "Confronto fra i campioni" (R11): metadati, test scelti e perché, alfa e beta per profiler
    (MetaPhlAn e Kraken separati), limiti. I numeri vengono da risultati/multicampione/ (motore/multicampione.R)."""
    m = d["multicampione"]
    md = m["metadati"]
    h = [f"<h3>{_('Confronto fra i campioni (analisi multicampione)')}</h3>",
         "<p>" + _("La diversità alfa dice quanti taxa ci sono dentro ogni campione; la diversità beta dice quanto i "
                   "campioni sono diversi fra loro (distanza di Aitchison sui dati trasformati con il logaritmo dei "
                   "rapporti, CLR, e distanza di Bray-Curtis sulle abbondanze relative). MetaPhlAn e Kraken restano "
                   "separati. Tabelle e grafici (in inglese) sono in {cartella}.",
                   cartella="<b>risultati/multicampione/</b>") + "</p>"]
    if not md["usati"]:
        h.append(riquadro(_("<b>Nessun test, solo descrizione.</b> Non è ancora scritto a quale gruppo appartiene ogni "
                            "campione. Nell'app, pagina Risultati, riquadro \"Analisi multicampione\": scrivi il gruppo "
                            "di ogni campione (ed eventualmente soggetto e tempo) e premi \"Salva e aggiorna l'analisi\". "
                            "I test si scelgono da soli in base al disegno. I dati restano nel file {file} della cartella "
                            "del progetto, che si può anche compilare con Excel.", file=f"<b>{md['file']}</b>")
                          + (" " + _("Il modello del file, con i nomi dei campioni, è stato creato adesso.")
                             if md["creato_ora"] else "")))
    for problema_ in md["problemi"]:
        h.append('<p class="piccolo"><b>' + _("Metadati:") + "</b> " + e(problema_) + "</p>")
    for profiler, nome in (("metaphlan", "MetaPhlAn"), ("kraken", "Kraken + Bracken")):
        p = m["profiler"].get(profiler)
        if not p:
            continue
        h.append(f"<h4>{nome}</h4>")
        if profiler == "kraken":
            h.append('<p class="piccolo">' + _("<b>Prototipo, non per decisioni</b> (come la sezione 3). Solo batteri, "
                                               "archei e funghi sopra la soglia di Kraken ({soglia} %).",
                                               soglia=pct(d["soglia"], 3)) + "</p>")
        else:
            h.append('<p class="piccolo">' + _("Tutte le specie, nessuna soglia di abbondanza.") + "</p>")
        h.append('<p class="piccolo">' + _("{n} campioni nell'analisi.", n=p["campioni"]) +
                 (" " + _("Esclusi perché senza taxa: {nomi}.", nomi=e(", ".join(p["esclusi"]))) if p["esclusi"] else "")
                 + "</p>")
        if p["errore"]:
            h.append(riquadro(e(p["errore"])))
        if not p["confronti"] and md["usati"]:
            h.append("<p>" + _("Nessun confronto: servono almeno 2 gruppi diversi fra i campioni di questo profilo.")
                     + "</p>")
        for c in p["confronti"]:
            h.append("<p>" + _("<b>Confronto per {fattore}</b> ({livelli}).",
                               fattore=e(_(multicampione.NOMI_FATTORI[c["fattore"]])),
                               livelli=e(", ".join(c["livelli"]))) + " " + e(c["motivo"]) + "</p>")
            alfa = [r for r in p["alfa"] if r["comparison"] == c["id"]]
            if alfa:
                h.append(tabella([_("Indice (alfa)"), _("Test"), "n", "p", _("Effetto")],
                                 [[e(_(INDICI.get(r["metric"], r["metric"]))), e(_(multicampione.TEST[r["test"]])),
                                   e(r["n"]), p_testo(r["p"]),
                                   (e(_(multicampione.EFFETTO.get(r["effect_size_name"], ""))) + " "
                                    + p_testo(r["effect_size"], 2)) if r["effect_size"] else e(nota_r(r["note"]))]
                                  for r in alfa], (2, 3), [17, 31, 7, 13, 32]))
                if c["test_alfa"] in ("kruskal_dunn", "friedman"):
                    h.append('<p class="piccolo">' + _("I confronti a coppie fra i gruppi sono nel file {f}.",
                                                       f=f"multicampione/{profiler}/alpha_posthoc.tsv") + "</p>")
            beta = [r for r in p["beta"] if r["comparison"] == c["id"]]
            righe_b, permutazioni = [], ""
            for dist in DISTANZE:   # la dispersione (PERMDISP) sta sempre nella stessa riga della PERMANOVA
                pa = next((r for r in beta if r["distance"] == dist and r["test"] == "PERMANOVA"), None)
                pd = next((r for r in beta if r["distance"] == dist and r["test"] == "PERMDISP"), None)
                if pa:
                    permutazioni = pa["permutations"]
                    righe_b.append([e(_(DISTANZE[dist])), f"R² {p_testo(pa['R2'], 2)}; p {p_testo(pa['p'])}"
                                    if pa["p"] else "", p_testo(pd["p"]) if pd else "", nota_dispersione(pa)])
            if righe_b:
                h.append(tabella([_("Distanza (beta)"), "PERMANOVA", _("PERMDISP, p"), _("Nota")], righe_b, (1, 2),
                                 [21, 24, 21, 34]))
                h.append('<p class="piccolo">' + _("PERMANOVA: quanto la composizione media differisce fra i gruppi "
                                                   "(R² = parte della variabilità spiegata dai gruppi). PERMDISP: se "
                                                   "i gruppi sono dispersi in modo diverso. {n} permutazioni {modo}, "
                                                   "seme {seme}.", n=e(permutazioni),
                                                   modo=_(multicampione.PERMUTA[c["permutazioni"]]),
                                                   seme=multicampione.SEME) + "</p>")
    h.append('<p class="piccolo">' + _(LIMITI_MULTI) + "</p>")
    h.append('<p class="piccolo">' + _("Metodi: zeri sostituiti con metà della più piccola abbondanza non nulla "
                                       "(sostituzione moltiplicativa, Martín-Fernández 2003) prima del CLR; PERMANOVA "
                                       "(Anderson 2001) sempre con PERMDISP (Anderson 2006), sulle stesse permutazioni; "
                                       "programmi e versioni in parametri.tsv.") + "</p>")
    return h


QC_UNA_PAGINA = 8    # campioni: fino a 8 la sezione QC sta intera in una pagina del PDF (misurato 22/10, test_qc.py)


def valori_qc(c):
    """[coppie prima, coppie dopo, Q30 prima, Q30 dopo, duplicati, lunghezza, ospite, coppie per i profili] di un
    campione di risultati.json "qc" (v0.4), già scritti per le persone; vuoto se il numero manca."""
    intero = lambda v: "" if v is None else lingua.migliaia(f"{v:,}")  # noqa: E731
    return [intero(c.get("coppie_prima")), intero(c.get("coppie_dopo")), pct(c.get("q30_prima_pct"), 1),
            pct(c.get("q30_dopo_pct"), 1), pct(c.get("duplicati_pct"), 1), pct(c.get("lunghezza_media"), 1),
            pct(c.get("ospite_pct"), 1), intero(c.get("coppie_profili"))]


def html_qc(d, n):
    """Sezione "Controllo qualità delle read" (v0.4): due tabelle di 5 colonne (il PDF di Qt rende male le tabelle
    larghe), una frase su come leggerle e il collegamento al report MultiQC."""
    qc = d["qc"]
    titolo = f"{n}. {_('Controllo qualità delle read')}"
    if not qc.get("campioni"):
        return [f"<h2>{titolo}</h2>", "<p>" + _("Controllo qualità non disponibile per questa analisi.") + "</p>"]
    # 22/10: nel PDF Qt spezza le tabelle fra le pagine (intestazione in fondo a una, riga del campione sull'altra) e non
    # rispetta né <thead> né una tabella esterna: rispetta solo il salto di pagina. La sezione parte su una pagina nuova
    # (fino a QC_UNA_PAGINA campioni ci sta intera); con più campioni anche la seconda tabella parte su una pagina nuova
    h = [f'<h2 style="page-break-before: always">{titolo}</h2>']
    singole = any(not c.get("a_coppie", True) for c in qc["campioni"])
    h.append("<p>" + _("Prima dei profili le read passano da fastp, che toglie le parti di bassa qualità e gli "
                       "adattatori, e poi da Bowtie 2, che toglie il DNA dell'ospite e dell'uomo.") + "</p>")
    righe = [[e(c["nome"])] + valori_qc(c) for c in qc["campioni"]]
    # Larghezze (22/10): nel PDF (510 punti di pagina) un numero non va a capo a metà fino a 8 cifre (34.000.000 coppie,
    # come C03_1; con 7 cifre "1.507.887" andava a capo) e ogni intestazione sta al massimo su 3 righe, "Ospite e uomo
    # tolti (%)" su 2: intestazioni a 8,5 punti e celle a 9 (dentro la cella: lo stile su <th> e <td> Qt lo ignora)
    def piccole(t):
        return re.sub(r"(<t([hd]) [^>]*>)(.*?)(</t[hd]>)",
                      lambda m: f'{m[1]}<span style="font-size:{"8.5" if m[2] == "h" else "9"}pt">{m[3]}</span>{m[4]}', t)
    h.append("<h3>" + _("Pulizia delle read (fastp)") + "</h3>")
    h.append(piccole(tabella([_("Campione"), _("Coppie prima"), _("Coppie dopo"), _("Basi ≥ Q30 prima (%)"),
                              _("Basi ≥ Q30 dopo (%)")], [r[:5] for r in righe], (1, 2, 3, 4), [22, 20, 20, 19, 19])))
    salto = ' style="page-break-before: always"' if len(righe) > QC_UNA_PAGINA else ""
    h.append(f"<h3{salto}>" + _("Duplicati, lunghezza e ospite") + "</h3>")
    h.append(piccole(tabella([_("Campione"), _("Duplicati, read grezze (%)"), _("Lunghezza media (basi)"),
                              _("Ospite e uomo tolti (%)"), _("Coppie per i profili")],
                             [[r[0]] + r[5:] for r in righe], (1, 2, 3, 4), [21, 18, 18, 24, 19])))
    if singole:
        h.append('<p class="piccolo">' + _("Per i campioni a lettura singola le \"coppie\" sono read singole.")
                 + "</p>")
    h.append('<p class="piccolo">' + _("<b>Come leggerle.</b> Q30 vuol dire al massimo 1 errore ogni 1.000 basi: più "
                                       "alta è la percentuale, migliori sono le read, e dopo fastp sale perché le "
                                       "parti peggiori sono state tolte. Molti duplicati vogliono dire molte copie "
                                       "della stessa molecola (poco DNA di partenza o molta PCR). I duplicati sono "
                                       "misurati sulle read grezze. La lunghezza media "
                                       "è quella dopo fastp. Più DNA dell'ospite si toglie, meno coppie restano per i "
                                       "profili: con poche coppie i taxa rari possono mancare (sezione 2).") + "</p>")
    if qc.get("multiqc"):
        f = qc["multiqc"]
        h.append("<p>" + _("Grafici e dettagli di fastp, FastQC e Bowtie 2 per ogni campione sono nel report MultiQC: "
                           "{link}", link=f'<a href="../{e(f)}">risultati/{e(f)}</a>') + "</p>")
    return h


def righe_parametri(d):
    """[(parametro, valore, cosa vuol dire, fonte)] (D022)."""
    p, prof = d["parametri"], profili_kraken()
    profilo = p.get("kraken_profile", "")
    fonte_k = _(prof.get(profilo, {}).get("fonte", ""))   # fonti di refs/kraken_profili.tsv, nel catalogo anche loro
    out = []
    if p.get("metaphlan_mode"):
        out.append((_("Modo di MetaPhlAn"), p["metaphlan_mode"],
                    _("fette = database diviso in 6 parti per stare in 16 GB; profilo delle specie uguale al modo "
                      "standard") if p["metaphlan_mode"] == "fette" else
                    _("standard = database intero in memoria (oltre 40 GB)"),
                    _("D016 (prova S2.1: 27 specie su 27 identiche); D051")))
    if p.get("metaphlan_index"):
        out.append((_("Database di MetaPhlAn"), p["metaphlan_index"], _("elenco dei geni marcatori delle specie"),
                    "Blanco-Miguez et al. 2023"))
    if p.get("kraken_db"):
        spiega = {"completo": _("PlusPF completo: batteri, archei, funghi, protozoi, uomo (consigliato)"),
                  "rapido": _("PlusPF-8, screening rapido: perde alcune specie fra 0,1 e 1 %")}.get(profilo, "")
        out.append((_("Database di Kraken"), f"{os.path.basename(p['kraken_db'])} ({profilo})", spiega,
                    _("D023 (prova K3); Wright, Comeau & Langille 2023")))

    def da_profilo(chiave_p, chiave_prof, nome, spiega):
        v = p.get(chiave_p)
        if v is None or v == "":
            return
        consigliato = prof.get(profilo, {}).get(chiave_prof)
        fonte = fonte_k if consigliato is not None and numero(consigliato) == numero(v) else \
            _("scelto nel progetto (consigliato per '{profilo}': {consigliato})", profilo=profilo,
              consigliato=lingua.virgola(consigliato))
        out.append((nome, lingua.virgola(v), spiega, fonte))   # 22/10: "0,1" con l'italiano, non "0.1"
    da_profilo("kraken_confidence", "confidence", _("Confidenza di Kraken (--confidence)"),
               _("quanto Kraken deve essere sicuro per assegnare una read (0 = accetta tutte)"))
    da_profilo("bracken_min_reads", "min_reads", _("Letture minime di Bracken (-t)"),
               _("sotto questo numero di read una specie non si stima"))
    da_profilo("abundance_threshold_pct", "soglia", _("Soglia di abbondanza di Kraken (%)"),
               _("sotto questa percentuale un taxon di Kraken non compare nelle tabelle del referto; MetaPhlAn non ha "
                 "soglia"))
    if p.get("bracken_readlen"):
        out.append((_("Lunghezza delle read per Bracken (-r)"), p["bracken_readlen"], _("lunghezza tipica delle read"),
                    _("impostazione della pipeline")))
    ospiti = [o for o in p.get("hosts", "").replace(",", " ").split() if o]
    if p.get("filter_human") in ("yes", "true", "True", "1"):
        ospiti.append("grch38")
    if ospiti:
        out.append((_("DNA tolto prima dell'analisi (ospite)"), ", ".join(_(OSPITI.get(o, o)) for o in ospiti),
                    _("genomi di riferimento usati per togliere il DNA dell'ospite e dell'uomo"),
                    _("D021 (stesse versioni dei profili pubblicati)")))
    return out


def dati_prodotto(richiesta):
    """La richiesta di sicurezza in parole: "pH 5,4, nitrito 60 ppm, ..., 30 giorni a 5 °C"."""
    prod = richiesta.get("prodotto", {})
    num = lambda v: lingua.virgola(f"{v:g}")  # noqa: E731
    parti = [_(NOMI_PRODOTTO[k]).format(num(v)) for k, v in prod.items() if k in NOMI_PRODOTTO]
    parti += [_("acido {acido} {ppm} ppm", acido=_(a), ppm=num(v)) for a, v in prod.get("acidi", {}).items()]
    if richiesta.get("profilo_temperatura"):
        parti.append(_("temperature: ") + "; ".join(_("{ore} ore a {t} °C", ore=num(o), t=num(t))
                                                    for t, o in richiesta["profilo_temperatura"]))
    else:
        parti.append(_("{giorni} giorni a {t} °C", giorni=num(richiesta.get('giorni', 30)),
                       t=num(prod.get('temperatura', 5))))
    for k, v in richiesta.get("conte", {}).items():
        parti.append(f"{_('conta dei batteri lattici') if k == 'lab_ufc_g' else _('conta totale')} "
                     f"{lingua.migliaia(f'{v:,.0f}')} ufc/g")
    return ", ".join(parti)


# ---------------- HTML (anche base del PDF) ----------------
def e(x):
    return html.escape(str(x))


def pct(v, cifre=2):
    return "" if v is None else lingua.virgola(f"{v:.{cifre}f}")


def tabella(intestazione, righe, allinea_destra=(), larghezze=None):
    larghezze = larghezze or [round(100 / len(intestazione))] * len(intestazione)
    t = ['<table border="1" cellspacing="0" cellpadding="4" width="100%">',
         "<tr>" + "".join(f'<th bgcolor="#dfe8f1" align="left" width="{w}%">{e(h)}</th>'
                          for h, w in zip(intestazione, larghezze)) + "</tr>"]
    for r in righe:
        t.append("<tr>" + "".join(f'<td align="{"right" if i in allinea_destra else "left"}">{c}</td>'
                                  for i, c in enumerate(r)) + "</tr>")
    t.append("</table>")
    return "\n".join(t)


# D115: nella prova con le persone Kraken e FSSP sono prototipi. Kraken dopo K4 (19/10: "Resta prototipo,
# testo nuovo"): tarato su comunità simulate e sullo Zymo, non ancora su salumi veri
PROTOTIPO_KRAKEN = N_("<b>Prototipo, non per decisioni.</b> Le impostazioni di Kraken sono tarate su comunità simulate e "
                      "su un campione di controllo (Zymo), non ancora su salumi veri: guarda questi numeri come un "
                      "confronto, non usarli per decidere.")
PROTOTIPO_FSSP = N_("<b>Prototipo, non per decisioni.</b> Il collegamento fra le specie trovate e i modelli FSSP è "
                    "ancora in prova: le previsioni non dicono se un prodotto è sicuro.")


# D141, D142 (20/10): i limiti si dichiarano con i numeri di L3 e K4. La prima frase usa i valori veri della corsa
LIMITE_KRAKEN = N_("Kraken mostra solo i taxa con almeno lo {soglia} % delle read assegnate da Bracken (non "
                   "classificate escluse; confidence {confidence}). Con i valori consigliati per il PlusPF completo "
                   "(confidence 0,5 e soglia 0,1 %), nelle comunità simulate le specie false sopra lo 0,1 % scendono "
                   "da 196 a 64 su 12 campioni, e nel campione di controllo Zymo da 10 a 0. I taxa sotto la soglia, "
                   "compresa <i>Listeria monocytogenes</i>, non compaiono nel profilo di Kraken: per la "
                   "<i>Listeria</i> serve un'analisi mirata.")
LIMITE_METAPHLAN = N_("MetaPhlAn non usa soglie, ma con molto DNA dell'ospite restano poche read dei taxa e i taxa rari "
                      "possono mancare. Nelle comunità simulate, <i>Listeria monocytogenes</i> allo 0,1 % è stata "
                      "trovata con circa 1.800 coppie di read sue e persa con circa 900; con il 95 % di ospite (circa "
                      "150.000 coppie di taxa) mancavano batteri fino allo 0,4 %.")


def limite_kraken(d):
    return _(LIMITE_KRAKEN, soglia=lingua.virgola(f"{d['soglia']:g}"),
             confidence=lingua.virgola(d["parametri"].get("kraken_confidence", "")))


def senza_html(testo):
    return re.sub(r"<[^>]+>", "", testo)


def riquadro(testo, colore="#fff4d6"):
    return f'<table border="0" cellpadding="8" width="100%"><tr><td bgcolor="{colore}">{testo}</td></tr></table>'


def html_referto(d):
    s = d["soglia"]
    h = ["<!DOCTYPE html>", f'<html lang="{lingua.attuale()}"><head><meta charset="utf-8">',
         "<title>" + _("Referto {progetto}", progetto=e(d['progetto'])) + "</title>",
         "<style>body{font-family:'DejaVu Sans',Arial,sans-serif;font-size:10pt;color:#222;max-width:60em;margin:auto}"
         "h1{color:#1f4e79}h2{color:#1f4e79;border-bottom:1px solid #1f4e79;margin-top:1.4em}h3{color:#333}"
         "table{border-collapse:collapse;margin:0.4em 0}td,th{border-color:#999}.piccolo{font-size:8.5pt;color:#555}"
         "</style></head><body>",
         "<h1>" + _("Referto dell'analisi: {progetto}", progetto=e(d['progetto'])) + "</h1>",
         _("<p>Corsa <b>{corsa}</b> &middot; motore profilemetananni {motore} &middot; {campioni}</p>",
           corsa=e(d['corsa']), motore=e(d['motore']),
           campioni=_("1 campione") if len(d['campioni']) == 1 else _("{n} campioni", n=len(d['campioni']))),
         riquadro(_("<b>Come leggere questo referto.</b> Le percentuali dicono quanto DNA di ogni taxon c'è nel "
                    "campione, non quante cellule vive per grammo. MetaPhlAn e Kraken sono due metodi diversi: i loro "
                    "risultati stanno in tabelle separate e non vanno sommati. In fondo trovi i parametri usati, gli "
                    "strumenti con le versioni e come citarli."), "#e8f1fb")]
    # in breve
    h.append(f"<h2>1. {_('In breve')}</h2>")
    esiti = {t["campione"]: t for t in d["triage"]}
    for c in d["campioni"]:
        h.append("<h3>" + _("Campione {nome}", nome=e(c['nome'])) + "</h3><ul>")
        if c["coppie"]:
            h.append(lingua.migliaia(_("<li>Coppie di read dopo aver tolto l'ospite: <b>{n}</b></li>",
                                       n=f"{c['coppie']:,}")))
        if c["nome"] in esiti:
            t = esiti[c["nome"]]
            h.append(_("<li>Controllo iniziale (triage): <b>{esito}</b>", esito=e(_(t['esito'])))
                     + (f" ({e('; '.join(t['motivi']))})" if t["motivi"] else "") + "</li>")
        for metodo in ("metaphlan", "kraken"):   # stessa frase del risultati.json (sintesi)
            if metodo in c["sintesi"]:
                h.append(f"<li>{e(c['sintesi'][metodo]['frase'])}</li>")
        h.append("</ul>")
    # MetaPhlAn
    h.append(f"<h2>2. {_('Specie trovate con MetaPhlAn')}</h2>")
    if any(c["metaphlan"] for c in d["campioni"]):
        h.append("<p>" + _(LIMITE_METAPHLAN) + "</p>")
    for c in d["campioni"]:
        if not c["metaphlan"]:
            continue
        somma = sum(v for _, v in c["metaphlan"]) or 1.0
        h.append("<h3>" + _("Campione {nome}", nome=e(c['nome'])) + "</h3>")
        h.append('<p class="piccolo">' + _("Read non assegnate a specie note (UNCLASSIFIED): {nc} %. {n} specie "
                                           "trovate, tutte qui sotto.", nc=pct(c["non_classificato"], 1),
                                           n=len(c['metaphlan'])) + "</p>")
        h.append(tabella([_("Specie"), _("% sul totale delle read"), _("% fra le specie trovate")],
                         [[f"<i>{e(n)}</i>", pct(v, 3), pct(100 * v / somma, 2)] for n, v in c["metaphlan"]], (1, 2),
                         [50, 25, 25]))
    # Kraken
    h.append(f"<h2>3. {_('Taxa trovati con Kraken + Bracken')}</h2>")
    h.append(riquadro(_(PROTOTIPO_KRAKEN)))
    if any(c["kraken"] for c in d["campioni"]):
        h.append("<p>" + limite_kraken(d) + "</p>")
    h.append('<p class="piccolo">' + _("Solo batteri, archei e funghi. Le read di mammifero rimaste dopo la pulizia, "
                                       "che il database assegna a <i>Homo sapiens</i> o a <i>Leishmania</i> perché non "
                                       "contiene il genoma del maiale, e i virus sono nella tabella successiva, a parte "
                                       "(D023).") + "</p>")
    for c in d["campioni"]:
        if not c["kraken"]:
            continue
        micro = [r for r in c["kraken"] if r["categoria"] in MICROBI]
        sopra = [r for r in micro if r["pct"] >= s]
        h.append("<h3>" + _("Campione {nome}", nome=e(c['nome'])) + "</h3>")
        h.append('<p class="piccolo">' + _("Taxa mostrati: {n} su {tot} (soglia di Kraken {soglia} %).", n=len(sopra),
                                           tot=len(micro), soglia=pct(s, 3)) + "</p>")
        h.append(tabella([_("Specie"), _("Gruppo"), _("% fra i taxa"), _("% sul totale assegnato")],
                         [[f"<i>{e(r['nome'])}</i>", e(_(r["categoria"])), pct(r["pct_microbi"], 2), pct(r["pct"], 3)]
                          for r in sopra], (2, 3), [44, 16, 20, 20]))
        altri = [r for r in c["kraken"] if r["categoria"] not in MICROBI]
        if altri:
            h.append(_("<p><b>Segnalati a parte (non del prodotto, o non validati)</b>, campione {nome}:</p>",
                       nome=e(c['nome'])))
            h.append(tabella([_("Nome"), _("Gruppo"), _("% sul totale assegnato"), _("Letture stimate")],
                             [[f"<i>{e(r['nome'])}</i>", e(_(r["categoria"])), pct(r["pct"], 3),
                               lingua.migliaia(f"{r['letture']:,}")] for r in altri], (2, 3), [46, 18, 18, 18]))
    h += html_diversita(d)
    # virus
    h.append(f"<h2>5. {_('Virus')}</h2>")
    h.append(riquadro(e(_(POSTILLA_VIRUS)) + "<br>" + (_("Richiesti nel progetto, ma non ancora calcolati da questa "
                                                          "versione del motore.") if d["virus"] else
                                                        _("Non richiesti in questo progetto."))))
    # sicurezza
    n = 6
    if d["sicurezza"]:
        h.append(f"<h2>{n}. {_('Sicurezza alimentare (previsioni di crescita)')}</h2>")
        h.append(riquadro(_(PROTOTIPO_FSSP)))
        h.append(riquadro(_("Previsioni da modelli di microbiologia previsionale (FSSP, DTU Food). Il profilo dice chi "
                            "c'è, non quanti per grammo: per la <i>Listeria</i> si calcolano scenari \"e se\" (1 e 100 "
                            "ufc/g). Una previsione non sostituisce le analisi di laboratorio.")))
        for r in d["sicurezza"]:
            h.append("<h3>" + e(_("Campione {nome}", nome=r['campione']) if r.get('campione') else _("Prodotto"))
                     + " (" + _("modalità {modo}", modo=e(_(r.get('modalita', '')))) + ")</h3>")
            if r.get("richiesta"):
                h.append('<p class="piccolo">' + _("Dati del prodotto: {dati}", dati=e(dati_prodotto(r["richiesta"])))
                         + "</p>")
            if r.get("profilo_usato"):
                h.append('<p class="piccolo">' + _("Profilo usato: {profilo}", profilo=e(r["profilo_usato"])) + "</p>")
            h.append(tabella([_("Scenario"), _("La Listeria può crescere?"), _("Giorni per aumentare di 100 volte"),
                              _("psi (1 o più = non cresce)")],
                             [[e(p["scenario"]), _("sì") if p["listeria_puo_crescere"] else _("no"),
                               pct(p.get("cento_volte_listeria_giorni"), 1) or _("non raggiunto"),
                               pct(p.get("psi_listeria"), 2) if p.get("psi_listeria") is not None else _("oltre un limite")]
                              for p in r["previsioni"]], (2, 3), [34, 18, 24, 24]))
            for a in r.get("avvisi", []):
                h.append(f'<p><b>{_("Attenzione")}</b> ({e(_(MODELLI.get(a.get("modello"), N_("modello"))))}): '
                         f'{e(a["messaggio"])}</p>')
            h.append("<ul>" + "".join(f"<li>{e(x)}</li>" for x in r.get("note", [])) + "</ul>")
        n += 1
    # controllo qualità (v0.4)
    h += html_qc(d, n)
    n += 1
    # parametri
    h.append(f"<h2>{n}. {_('Parametri usati')}</h2>")
    h.append(tabella([_("Parametro"), _("Valore"), _("Cosa vuol dire"), _("Fonte")],
                     [[e(a), e(b), e(c), f'<span class="piccolo">{e(f)}</span>'] for a, b, c, f in righe_parametri(d)],
                     larghezze=[20, 20, 30, 30]))
    n += 1
    # strumenti
    h.append(f"<h2>{n}. {_('Strumenti e database usati')}</h2>")
    h.append(tabella([_("Strumento o database"), _("Versione")],
                     [[e(_(VERSIONI.get(k, k))), e(v)] for k, v in d["versioni"].items() if v], larghezze=[45, 55]))
    citazioni = d["citazioni"] + [r["citazione"] for r in d["sicurezza"][:1]] + citazioni_multi(d)
    h.append(f"<p><b>{_('Come citarli:')}</b></p><ol>" + "".join(f"<li>{e(c)}</li>" for c in citazioni) + "</ol>")
    n += 1
    h.append(f"<h2>{n}. {_('Parole usate')}</h2><ul>" + "".join(f"<li><b>{e(_(a))}</b>: {e(_(b))}</li>"
                                                              for a, b in GLOSSARIO) + "</ul>")
    h.append("</body></html>")
    return "\n".join(h)


# ---------------- Excel ----------------
def cella(v):
    # niente formule nascoste nei testi (un nome che inizia con = + - @ resta testo)
    return "'" + v if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") else v


def excel(d, percorso):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    fogli = []

    def foglio(nome, intest, righe, nota=None):
        ws = wb.create_sheet(nome)
        if nota:                                   # D115: la scritta "prototipo" sopra le colonne
            ws.append([nota])
            ws["A1"].font = Font(bold=True, color="9C5700")
        ws.append(intest)
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
        for r in righe:
            ws.append([cella(v) for v in r])
        ws.freeze_panes = "A3" if nota else "A2"
        for i, h in enumerate(intest):
            ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = max(12, min(60, len(h) + 4))
        fogli.append(nome)
    leggimi = wb.active
    leggimi.title = _("Leggimi")
    kraken = any(c["kraken"] for c in d["campioni"])   # 22/10: soglia e limite di Kraken solo se Kraken ha girato
    for r in [[_("Referto dell'analisi: {progetto}", progetto=d['progetto'])],
              [_("Corsa {corsa}, motore {motore}", corsa=d['corsa'], motore=d['motore'])], [],
              [_("MetaPhlAn e Kraken sono metodi diversi: stanno in fogli separati e non vanno sommati.")],
              [_("Le percentuali sono DNA relativo, non cellule vive per grammo.")],
              *([[_("Soglia di abbondanza (solo Kraken): {soglia} % (qui ci sono anche i taxa sotto soglia). "
                    "MetaPhlAn non ha soglia.", soglia=lingua.virgola(f"{d['soglia']:g}"))],
                  [senza_html(limite_kraken(d))]] if kraken else []),
              [senza_html(_(LIMITE_METAPHLAN))],
              [_(POSTILLA_VIRUS)]]:
        leggimi.append(r)
    leggimi.column_dimensions["A"].width = 110
    si, no = _("sì"), _("no")
    proto_k = senza_html(_(PROTOTIPO_KRAKEN))
    foglio(_("MetaPhlAn_specie"), [_("Campione"), _("Specie"), _("% sul totale delle read"),
                                   _("% fra le specie trovate")],
           [[c["nome"], n, v, 100 * v / (sum(x for _x, x in c["metaphlan"]) or 1)]
            for c in d["campioni"] if c["metaphlan"] for n, v in c["metaphlan"]])
    foglio(_("Kraken_taxa"), [_("Campione"), _("Specie"), _("Gruppo"), _("TaxID NCBI"), _("Letture stimate"),
                              _("% sul totale assegnato"), _("% fra i taxa"), _("Sopra soglia")],
           [[c["nome"], r["nome"], _(r["categoria"]), r["taxid"], r["letture"], r["pct"], r["pct_microbi"],
             si if r["pct"] >= d["soglia"] else no]
            for c in d["campioni"] if c["kraken"] for r in c["kraken"] if r["categoria"] in MICROBI], proto_k)
    foglio(_("Kraken_a_parte"), [_("Campione"), _("Nome"), _("Gruppo"), _("TaxID NCBI"), _("Letture stimate"),
                                 _("% sul totale assegnato")],
           [[c["nome"], r["nome"], _(r["categoria"]), r["taxid"], r["letture"], r["pct"]]
            for c in d["campioni"] if c["kraken"] for r in c["kraken"] if r["categoria"] not in MICROBI], proto_k)
    for metodo, nome, colonne in (("metaphlan", _("Diversità_MetaPhlAn"), [_("Read del campione"), _("Read usate")]),
                                  ("kraken", _("Diversità_Kraken"), [_("Letture dei taxa"), _("Letture usate")])):
        if d["diversita"][metodo]:
            foglio(nome, [_("Campione"), *colonne, _("Stato"),
                          _("Specie (ricchezza)") if metodo == "metaphlan" else _("Taxa (ricchezza)"), "Shannon", "Simpson"],
                   [[r["campione"], r["letture"], r["letture_usate"], _(STATO_DIVERSITA[r["stato"]]), r["ricchezza"],
                     r["shannon"], r["simpson"]] for r in d["diversita"][metodo]],
                   proto_k if metodo == "kraken" else None)
    multi = d.get("multicampione") or {}
    for metodo, nome in (("metaphlan", _("Multi_MetaPhlAn")), ("kraken", _("Multi_Kraken"))):   # mai insieme
        p = multi.get("profiler", {}).get(metodo) if multi.get("attiva") else None
        if p and (p["alfa"] or p["beta"]):
            foglio(nome, [_("Confronto"), _("Diversità"), _("Indice o distanza"), _("Test"), "n", _("Statistica"),
                          _("Valore"), "p", _("Effetto"), _("Valore dell'effetto"), _("Permutazioni"), _("Nota")],
                   righe_multi_excel(p), proto_k if metodo == "kraken" else None)
    if d["triage"]:
        foglio("Triage", [_("Campione"), _("Esito"), _("Motivi"), _("Punteggio")],
               [[t["campione"], _(t["esito"]), "; ".join(t["motivi"]), t["punteggio"]] for t in d["triage"]])
    if d["sicurezza"]:
        foglio(_("Sicurezza"), [_("Campione"), _("Modalità"), _("Scenario"), _("Modello"), _("Listeria può crescere"),
                                _("psi Listeria"), _("Giorni per 100 volte (Listeria)"), _("Giorni per 100 volte (LAB)")],
               [[r.get("campione") or _("prodotto"), _(r.get("modalita")) if r.get("modalita") else r.get("modalita"),
                 p["scenario"], p["modello"], si if p["listeria_puo_crescere"] else no, p.get("psi_listeria"),
                 p.get("cento_volte_listeria_giorni"), p.get("cento_volte_lab_giorni")]
                for r in d["sicurezza"] for p in r["previsioni"]], senza_html(_(PROTOTIPO_FSSP)))
    if d["qc"].get("campioni"):
        foglio("QC", [_("Campione"), _("Coppie prima di fastp"), _("Coppie dopo fastp"), _("% basi ≥ Q30 prima"),
                      _("% basi ≥ Q30 dopo"), _("% duplicati (read grezze)"), _("Lunghezza media dopo fastp (basi)"),
                      _("% ospite e uomo tolti"), _("Coppie per i profili")],
               [[c["nome"], c.get("coppie_prima"), c.get("coppie_dopo"), c.get("q30_prima_pct"), c.get("q30_dopo_pct"),
                 c.get("duplicati_pct"), c.get("lunghezza_media"), c.get("ospite_pct"), c.get("coppie_profili")]
                for c in d["qc"]["campioni"]])
    foglio(_("Parametri"), [_("Parametro"), _("Valore"), _("Cosa vuol dire"), _("Fonte")], righe_parametri(d))
    citazioni = d["citazioni"] + [r["citazione"] for r in d["sicurezza"][:1]] + citazioni_multi(d)
    foglio(_("Strumenti_citazioni"), [_("Strumento o database"), _("Versione")],
           [[_(VERSIONI.get(k, k)), v] for k, v in d["versioni"].items()] + [[]] +
           [[_("Citazione"), c] for c in citazioni])
    wb.save(percorso)
    return fogli


# ---------------- PDF ----------------
def pdf(testo_html, percorso):
    """PDF dallo stesso HTML con QTextDocument (PySide6 di pmn_reports_py), senza schermo."""
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtCore import QMarginsF
    from PySide6.QtGui import QGuiApplication, QPageLayout, QPageSize, QPdfWriter, QTextDocument
    app = QGuiApplication.instance() or QGuiApplication(["referto"])  # noqa: F841 (serve vivo fino alla fine)
    w = QPdfWriter(str(percorso))
    w.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    w.setPageMargins(QMarginsF(15, 15, 15, 15), QPageLayout.Unit.Millimeter)
    w.setTitle(_("Referto profilemetananni"))
    w.setCreator("profilemetananni")
    doc = QTextDocument()
    doc.setHtml(testo_html)
    doc.print_(w)
    del w
    dati = Path(percorso).read_bytes() if Path(percorso).is_file() else b""
    if dati[:4] != b"%PDF":
        raise RuntimeError("PDF non scritto")
    return len(re.findall(rb"/Type\s*/Page(?!s)", dati))


def problema(file, modulo, ex):
    """Perché un file del referto manca, in italiano semplice (v0.3); il dettaglio tecnico va nel log."""
    if isinstance(ex, ImportError):
        return _("manca {file}: nel motore non c'è il modulo che lo scrive ({modulo})", file=_(file), modulo=modulo)
    if isinstance(ex, PermissionError):
        return _("manca {file}: il file è aperto in un altro programma o la cartella non si può scrivere", file=_(file))
    return _("manca {file}: errore imprevisto (il dettaglio è nel registro del motore)", file=_(file))


def avviso_parziale(problemi, c_e_la_pagina):
    """Testo dell'avviso REFERTO_PARZIALE (v0.3), uguale per il motore vero e il finto."""
    if not c_e_la_pagina:
        problemi = list(problemi) + [_("manca la pagina del referto")]
    testo = _("Il referto non è completo: ") + "; ".join(problemi) + "."
    return testo + (_(" La pagina del referto si apre lo stesso.") if c_e_la_pagina else "")


def main():
    ap = argparse.ArgumentParser(description="Referto HTML + Excel + PDF dei risultati di pmn-motore.")
    ap.add_argument("--risultati", required=True)
    ap.add_argument("--progetto")
    a = ap.parse_args()
    risultati = Path(a.risultati).resolve()
    if not (risultati / "risultati.json").is_file():
        print(json.dumps({"errore": f"manca {risultati / 'risultati.json'}"}, ensure_ascii=False))
        return 1
    cfg = {}
    if a.progetto and Path(a.progetto).is_file():
        try:
            import yaml
            cfg = yaml.safe_load(Path(a.progetto).read_text(encoding="utf-8")) or {}
        except Exception:
            cfg = {}
        cfg = cfg if isinstance(cfg, dict) else {}
    d = raccogli(risultati, cfg)
    problemi = []
    try:   # R11: con almeno 2 campioni; un errore qui non ferma il referto
        d["multicampione"] = multicampione.esegui(risultati, d)
    except Exception as ex:
        d["multicampione"] = {"attiva": False}
        problemi.append(_("manca l'analisi multicampione: errore imprevisto (il dettaglio è nel registro del motore)"))
        print(f"multicampione: {type(ex).__name__}: {ex}", file=sys.stderr)
    cartella = risultati / "referto"
    cartella.mkdir(exist_ok=True)
    for vecchio in ("referto.html", "referto.xlsx", "referto.pdf"):   # niente file di un referto precedente
        (cartella / vecchio).unlink(missing_ok=True)
    testo = html_referto(d)
    uscita = {"html": str(cartella / "referto.html"), "xlsx": None, "pdf": None, "problemi": problemi}
    (cartella / "referto.html").write_text(testo, encoding="utf-8")
    uscita["dettagli"] = []   # testo tecnico per il registro (log); per l'utente solo "problemi", in italiano
    try:
        excel(d, cartella / "referto.xlsx")
        uscita["xlsx"] = str(cartella / "referto.xlsx")
    except Exception as ex:   # manca openpyxl o il file non si scrive: il resto del referto c'e' lo stesso
        uscita["problemi"].append(problema(N_("l'Excel"), "openpyxl", ex))
        uscita["dettagli"].append(f"Excel: {type(ex).__name__}: {ex}")
    try:
        uscita["pagine_pdf"] = pdf(testo, cartella / "referto.pdf")
        uscita["pdf"] = str(cartella / "referto.pdf")
    except Exception as ex:
        uscita["problemi"].append(problema(N_("il PDF"), "PySide6", ex))
        uscita["dettagli"].append(f"PDF: {type(ex).__name__}: {ex}")
    print(json.dumps(uscita, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
