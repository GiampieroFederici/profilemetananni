# -*- coding: utf-8 -*-
"""Fotografia di tutti i testi che l'utente vede (sessione Lingua IT/EN, 18/10/2026).

Gira l'app con il finto motore attraverso le sue pagine e i suoi casi, poi il finto motore da solo (eventi, referto
HTML ed Excel) e alcune funzioni del motore vero che scrivono testi. Tutto quello che si legge finisce in un
dizionario {caso: [testi]}, con le parti che cambiano da una volta all'altra (date, ore, cartelle temporanee) già
sostituite da segnaposto.

  python tests/lingua_foto.py --scrivi    fotografia dell'italiano di oggi in tests/dati/lingua_it.json
                                          (fatta sul codice di f567f8e, PRIMA della lingua: è il riferimento)
  tests/test_lingua.py                    rifà la fotografia e la confronta: con l'italiano deve essere identica

Gira con l'ambiente dell'app su Windows:  app\\.venv\\Scripts\\python tests/lingua_foto.py --scrivi
"""
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

RADICE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RADICE))
sys.path.insert(0, str(RADICE / "motore"))
FOTO_IT = RADICE / "tests" / "dati" / "lingua_it.json"
FINTO = RADICE / "installer" / "finto_motore.py"
ATTESA = 60


# ---------------- segnaposto per le parti che cambiano ----------------
def pulisci(testo, cartelle=()):
    t = str(testo)
    for c in sorted(cartelle, key=lambda c: -len(str(c))):
        for forma in {str(c), str(c).replace("\\", "/"), Path(c).as_posix()}:
            t = t.replace(forma, "<CARTELLA>")
    t = re.sub(r"\d{4}-\d{2}-\d{2}T\d{2}[-:]\d{2}[-:]\d{2}(?:[+-]\d{2}:\d{2})?", "<ORA>", t)
    t = re.sub(r"\d{2}/\d{2}/\d{4}, \d{2}:\d{2}", "<DATA>", t)
    t = re.sub(r"progetto_\d{4}-\d{2}-\d{2}", "progetto_<OGGI>", t)
    t = re.sub(r"<CARTELLA>[\\/][^\s\"<>]*", "<CARTELLA>/...", t)
    return t


# ---------------- app ----------------
def testi_widget(radice, salta=("lingua",)):
    """Tutti i testi di una finestra o pagina, nell'ordine in cui Qt ha creato i widget."""
    from PySide6.QtCore import QCoreApplication, QEvent
    from PySide6.QtWidgets import (QAbstractButton, QComboBox, QDialog, QLabel, QLineEdit, QPlainTextEdit, QTabBar,
                                   QWidget)
    QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
    QCoreApplication.processEvents()
    out = []
    for w in [radice] + radice.findChildren(QWidget):
        if w.objectName() in salta or any(p.objectName() in salta for p in _antenati(w)):
            continue
        if isinstance(w, QDialog) or (w.isWindow() and w is not radice):
            continue
        parti = []
        if isinstance(w, (QLabel, QAbstractButton)) and w.text():
            parti.append(w.text())
        if isinstance(w, QLineEdit) and w.placeholderText():
            parti.append("[segnaposto] " + w.placeholderText())
        if isinstance(w, QComboBox):
            parti += ["[voce] " + w.itemText(i) for i in range(w.count())]
        if isinstance(w, QTabBar):
            parti += ["[scheda] " + w.tabText(i) for i in range(w.count())]
        if isinstance(w, QPlainTextEdit) and w.isReadOnly() and w.toPlainText():
            parti.append("[file] " + w.toPlainText())
        for tipo, valore in (("[suggerimento]", w.toolTip()), ("[nome]", w.accessibleName()),
                             ("[descrizione]", w.accessibleDescription())):
            if valore:
                parti.append(f"{tipo} {valore}")
        if w.isWindow() and w.windowTitle():
            parti.append("[titolo] " + w.windowTitle())
        out += parti
    return out


def _antenati(w):
    p = w.parentWidget()
    while p is not None:
        yield p
        p = p.parentWidget()


def aspetta(condizione, secondi=ATTESA):
    from PySide6.QtWidgets import QApplication
    fine = time.monotonic() + secondi
    while not condizione():
        if time.monotonic() > fine:
            raise TimeoutError("il finto motore non ha finito in tempo")
        QApplication.processEvents()
        time.sleep(0.01)
    QApplication.processEvents()


def foto_app(lingua_scelta=None):
    """{caso: [testi]} dell'app con il finto motore. lingua_scelta: None = quella predefinita (l'italiano)."""
    from PySide6.QtWidgets import QApplication, QPushButton
    from app import finestra, main, progetto
    try:                                  # sul codice di prima della lingua (nuova fotografia per l'unione) non c'è
        from app.lingua import _ as tr
    except ImportError:
        def tr(testo):
            return testo
    from app.pagine import esecuzione, scelte, sicurezza
    from app.pagine import motore as pagina_motore
    from app.pagine.base import numero

    app = QApplication.instance() or QApplication([])
    main.prepara(app)
    tmp = Path(tempfile.mkdtemp(prefix="pmn_lingua_"))
    vecchi = {k: os.environ.get(k) for k in ("PMN_MOTORE", "PMN_FINTO_OPZIONI", "TEMP", "TMP")}
    cartella_progetti, avvio_min = progetto.cartella_progetti, finestra.AVVIO_MIN_MS
    spazio_windows = pagina_motore.spazio_windows_gb
    foto, cartelle = {}, [tmp]
    os.environ.update(PMN_MOTORE="finto", PMN_FINTO_OPZIONI="--velocita 0")
    (tmp / "temp").mkdir()
    os.environ["TEMP"] = os.environ["TMP"] = str(tmp / "temp")
    radice = tmp / "Documenti" / "ProfileMetaNanni"
    progetto.cartella_progetti = lambda: radice
    finestra.AVVIO_MIN_MS = 0
    pagina_motore.spazio_windows_gb = lambda: None     # B2 (P8): la fotografia non dipende dal disco di questo PC
    finestre = []

    def nuova():
        if lingua_scelta:                    # prima di creare le pagine: i testi si scrivono quando nascono
            from app import lingua
            lingua.imposta(lingua_scelta)
        w = finestra.Finestra(movimento=False)
        finestre.append(w)
        w.show()
        w.avvia()
        aspetta(lambda: w.attuale == "home")
        return w

    def scatta(caso, w, *altro):
        foto[caso] = [pulisci(t, cartelle) for t in testi_widget(w) + [x for a in altro for x in a]]

    def scenario(nome=None):
        os.environ["PMN_FINTO_OPZIONI"] = "--velocita 0" + (f" --scenario {nome}" if nome else "")

    def premi(w):
        w.principale.click()
        QApplication.processEvents()

    def stima_finita(w):
        aspetta(lambda: not w.pagine["scelte"].client.in_corso())

    try:
        # 1. avvio e progetti
        w = nuova()
        scatta("01 home senza progetti", w)
        w.pagine["home"].nome.setText("nome con spazi")
        premi(w)
        scatta("02 home nome sbagliato", w)
        w.pagine["home"].nome.setText("")
        w.vai("home")
        premi(w)                                     # crea il progetto proposto, va a Campioni
        scatta("03 campioni nuovo", w)
        c = w.pagine["campioni"]
        premi(w)
        scatta("04 campioni senza codici", w)
        c.testo_codici.setPlainText("SRR4875279\nciao\nERR1\nx1\nx2\nx3\nx4\nx5")
        scatta("05 campioni codici", w)
        c.scelta_pc.pulsante.setChecked(True)
        premi(w)
        scatta("06 campioni senza cartella", w)
        vuota = tmp / "vuota"
        vuota.mkdir()
        c.usa_cartella(vuota)
        scatta("07 cartella vuota", w)
        brutta = tmp / "a;b"
        brutta.mkdir()
        (brutta / "x_R1.fastq.gz").write_bytes(b"")
        c.usa_cartella(brutta)
        scatta("08 cartella brutta", w)
        reads = tmp / "reads"
        reads.mkdir()
        for n in ("A_R1.fastq.gz", "A_R2.fastq.gz", "B_1.fq.gz", "B_2.fq.gz", "solo.fastq.gz", "note.txt",
                  "orfano_R1.fastq.gz"):
            (reads / n).write_bytes(b"")
        c.usa_cartella(reads)
        scatta("09 cartella con read", w)
        uno = tmp / "uno"
        uno.mkdir()
        (uno / "Z_R1.fastq.gz").write_bytes(b"")
        (uno / "Z_R2.fastq.gz").write_bytes(b"")
        c.usa_cartella(uno)
        scatta("10 cartella con un campione", w)
        for chiave in ("maiale", "bovino", "pecora", "altro"):
            c.prodotto.setCurrentIndex(c.prodotto.findData(chiave))
            scatta(f"11 prodotto {chiave}", w)
        c.prodotto.setCurrentIndex(c.prodotto.findData("maiale"))
        c.scelta_ncbi.pulsante.setChecked(True)
        c.testo_codici.setPlainText("SRR4875279")
        premi(w)                                     # Scelte, con la stima
        stima_finita(w)
        scatta("12 scelte stima", w)
        s = w.pagine["scelte"]
        s.apri_esperti.setChecked(True)
        s.apri_file.setChecked(True)
        scatta("13 scelte esperti aperti", w)
        s.rapido.pulsante.setChecked(True)
        stima_finita(w)
        scatta("14 scelte screening rapido", w)
        s.completo.pulsante.setChecked(True)
        stima_finita(w)
        scenario("TEMPO_LUNGO")
        s.virus.setChecked(True)
        stima_finita(w)
        scatta("15 scelte tempo lungo", w)
        scenario()
        s.virus.setChecked(False)
        stima_finita(w)
        s._errore({"codice": "DB_MANCANTE", "messaggio": "Manca il database di Kraken.",
                   "cosa_fare": "Apri la pagina Motore e premi Installa o completa."})
        scatta("16 scelte errore di stima", w)
        s._errore({"codice": "SCONOSCIUTO", "messaggio": "X.", "cosa_fare": "Y."})
        scatta("17 scelte errore sconosciuto", w)
        w.vai("home")
        scatta("18 home con un progetto", w)

        # 2. analisi: ogni scenario del finto motore, un progetto nuovo ciascuno
        for nome in ("ok", "DB_MANCANTE", "READ_NON_TROVATE", "annullata", "GIA_IN_CORSO", "CAMPIONE_DA_CONTROLLARE",
                     "TEMPO_LUNGO"):
            scenario(None if nome == "ok" else nome)
            p = progetto.nuovo(f"corsa_{nome.lower()}", radice)
            p.scrivi({"project_name": p.nome, "motore": {"metaphlan_modo": "fette", "kraken_db": "completo"},
                      "app": {"prodotto": "maiale", "scelte": True}})
            w.contesto.progetto = p
            w.vai("esecuzione")
            scatta(f"20 esecuzione {nome} prima", w)
            es = w.pagine["esecuzione"]
            visti = []
            es.client.evento.connect(lambda e, es=es, visti=visti: visti.append(
                f"{es.messaggio.text()} | {es.pct.text()} | {es.eta.text()} | {es.eta_nota.text()}"))
            w.contesto.avvia_analisi = True
            w.vai("esecuzione")
            aspetta(lambda: es.stato not in ("in_corso", "ferma_chiesto"))
            es.dettagli.setChecked(True)
            scatta(f"21 esecuzione {nome} fine", w, visti)
            es.dettagli.setChecked(False)
        scenario()
        es = w.pagine["esecuzione"]                  # orologio locale: passo che sfora, ultimo passo
        es.stato = "in_corso"
        es.ultimo = (time.monotonic() - 100, 10, 50, 20)
        es._scorre()
        uno = [es.eta.text(), es.eta_nota.text()]
        es.ultimo = (time.monotonic() - 100, 10, 50, None)
        es._scorre()
        es.stato = "vuota"
        scatta("22 esecuzione orologio", w, uno, [es.eta.text()])
        es._passi("ospite", fermo=True)
        scatta("23 esecuzione passi fermi", w)
        es._passi("kraken", errore=True)
        scatta("24 esecuzione passi errore", w)

        # 3. risultati (progetto "corsa_ok")
        p = progetto.Progetto(radice / "corsa_ok")
        w.contesto.progetto = p
        w.vai("risultati")
        r = w.pagine["risultati"]
        for i in range(3):
            r.linguette.setCurrentIndex(i)
            scatta(f"30 risultati scheda {i}", w)
        dom = [{"nome": "Latilactobacillus sakei", "pct": 24.5, "pct_microbi": 40.4}]
        cc = [{"nome": "mini", "sintesi": {"metaphlan": {"non_assegnata_pct": 56.2, "specie": 1, "dominanti": dom},
                                           "kraken": {"microbi": 1, "dominanti": dom}}},
              {"nome": "due", "sintesi": {}}]
        scatta("31 risultati poche specie", r._scheda_metaphlan(cc), testi_widget(r._scheda_kraken(cc)))
        r.apri_referto()
        foto["32 referto nella finestra"] = [pulisci(r.finestra_referto.windowTitle(), cartelle)] + [
            pulisci(t, cartelle) for t in testi_widget(r.finestra_referto)]
        r.finestra_referto.close()
        pd = progetto.Progetto(radice / "corsa_campione_da_controllare")
        w.contesto.progetto = pd
        w.vai("risultati")
        scatta("33 risultati da controllare", w)
        w.contesto.progetto = progetto.Progetto(radice / "corsa_db_mancante")
        w.vai("risultati")
        scatta("34 risultati senza risultati", w)
        w.contesto.progetto = None
        for chiave in ("campioni", "scelte", "esecuzione", "risultati", "sicurezza"):
            w.vai(chiave)
            scatta(f"35 {chiave} senza progetto", w)

        # 4. sicurezza alimentare
        w.contesto.progetto = progetto.Progetto(radice / "corsa_db_mancante")
        w.vai("sicurezza")
        scatta("40 sicurezza senza risultati", w)
        w.contesto.progetto = p
        w.vai("sicurezza")
        sic = w.pagine["sicurezza"]
        sic.altro.setChecked(True)
        scatta("41 sicurezza vuota", w)
        sic.calcola.click()
        scatta("42 sicurezza manca la temperatura", w)
        sic.temperatura.setText("cinque")
        sic.calcola.click()
        scatta("43 sicurezza non un numero", w)
        casi = [("44 salsiccia cresce", dict(temperatura="5", giorni="21", sale="3,5", ph="6,0", nitrito="100"),
                 False, {}),
                ("45 con i lattici", dict(temperatura="5", giorni="21", sale="3,5", ph="6,0", nitrito="",
                                          lab="1000000"), False, {}),
                ("46 salame fuori campo", dict(temperatura="15", giorni="60", sale="0,90", ph="5,2", nitrito="100",
                                                lab=""), True, {}),
                ("47 tutto e come fssp", dict(temperatura="8", giorni="30", sale="4", ph="5,8", nitrito="60",
                                               co2="20", fenoli="5"), False, {"lattico": "7000", "acetico": "500"}),
                ("48 fuori dai limiti", dict(temperatura="80", giorni="21", sale="3,5", ph="6,0", nitrito="",
                                              co2="", fenoli=""), False, {"lattico": "", "acetico": ""})]
        for caso, campi, aw, acidi in casi:
            sic.con_aw.setChecked(aw)
            sic.con_sale.setChecked(not aw)
            for k, v in campi.items():
                getattr(sic, k).setText(v)
            for k, v in acidi.items():
                sic.acidi[k].setText(v)
            if caso.startswith("47"):
                sic.come_fssp.setChecked(True)
                sic.lag.setChecked(True)
            sic.calcola.click()
            aspetta(lambda: not sic.client.in_corso() and sic.calcola.isEnabled())
            scatta(f"{caso}", w, [n for n, *_ in getattr(sic, "grafico", None).serie] if hasattr(sic, "grafico")
                   else [])
        sic.migliorata.setChecked(True)
        sic.lag.setChecked(False)
        for k in ("lattico", "acetico"):
            sic.acidi[k].setText("")
        for k in ("co2", "fenoli"):
            getattr(sic, k).setText("")
        for k, v in dict(temperatura="5", giorni="21", sale="3,5", ph="6,0", nitrito="100", lab="").items():
            getattr(sic, k).setText(v)
        sic.con_sale.setChecked(True)
        sic.calcola.click()
        aspetta(lambda: not sic.client.in_corso() and sic.calcola.isEnabled())
        premi(w)                                     # aggiungi al referto
        aspetta(lambda: w.attuale == "risultati")
        scatta("49 risultati con la sicurezza", w)
        r.apri_referto()
        foto["49b referto con la sicurezza"] = [pulisci(t, cartelle) for t in testi_widget(r.finestra_referto)]
        r.finestra_referto.close()

        # 5. motore e crediti
        w.vai("motore")
        scatta("50 motore", w)
        m = w.pagine["motore"]
        m.verifica()
        aspetta(lambda: not m.client.in_corso())
        aspetta(lambda: not m.client_info.in_corso())
        scatta("51 motore verifica verde", w)
        scenario("DB_MANCANTE")
        m.verifica()
        aspetta(lambda: not m.client.in_corso())
        aspetta(lambda: not m.client_info.in_corso())
        scatta("52 motore verifica rossa", w)
        scenario()
        m.installa()
        QApplication.processEvents()
        scatta("53 motore installa in corso", w)
        aspetta(lambda: not m.client.in_corso())
        aspetta(lambda: not m.client_info.in_corso())
        scatta("54 motore installato", w)
        scenario("annullata")
        m.installa()
        aspetta(lambda: not m.client.in_corso())
        aspetta(lambda: not m.client_info.in_corso())
        scatta("55 motore installa annullata", w)
        scenario()
        info = w.contesto.info
        w.contesto.info = dict(info, finto=False, database=info["database"][:1] + [
            dict(info["database"][1], installato=False), {"nome": "ignoto", "installato": True}],
            consigli=info["consigli"] + [{"codice": "RAM_INSUFFICIENTE", "livello": "avviso", "messaggio": "Poca RAM.",
                                          "fonte": "x"}])
        m._disegna()
        w._stato_motore()
        scatta("56 motore manca qualcosa", w)
        w.contesto.info = None
        m._disegna({"codice": "MOTORE_NON_AVVIATO"})
        w._stato_motore()
        scatta("57 motore non risponde", w)
        m._disegna({"codice": "ALTRO", "messaggio": "Messaggio.", "cosa_fare": "Cosa fare."})
        w.vai("home")
        scatta("58 home senza motore", w)
        w.contesto.info = info
        w._stato_motore()
        w.vai("crediti")
        scatta("59 crediti", w)
        w.pagine["crediti"].copia()
        scatta("59b crediti copiate", w)
        w.contesto.progetto_in_analisi = p
        w.vai("home")
        w.pagine["home"].occupato()
        w.pagine["home"].occupato(progetto.Progetto(radice / "corsa_annullata"))
        scatta("60 home analisi in corso", w)
        w.contesto.progetto_in_analisi = None

        # 6. campione d'esempio, avvio lento, motore che non parte
        w.vai("home")
        prova = tr("Prova il campione d'esempio")
        next(b for b in w.pagine["home"].findChildren(QPushButton) if b.text() == prova).click()
        stima_finita(w)
        scatta("61 esempio scelte", w)
        w.vai("campioni")
        scatta("62 esempio campioni", w)
        foto["63 avvio"] = [pulisci(w.messaggio_avvio.text(), cartelle), tr(finestra.AVVIO_LENTO)]
        from app import motore as client
        vecchio = client.comando_motore
        client.comando_motore = lambda: (["programma_che_non_esiste_pmn.exe"], [])
        try:
            w2 = nuova()
            scatta("64 motore che non parte", w2)
            w2.vai("motore")
            aspetta(lambda: not w2.pagine["motore"].client_info.in_corso())
            scatta("65 motore che non parte, pagina motore", w2)
        finally:
            client.comando_motore = vecchio

        # 7. numeri e durate (solo formato)
        valori = [0, 0.4, 1, 59, 89.6, 90, 3599, 5399, 5400, 9.1 * 3600, 123456.789]
        foto["70 formati"] = ([scelte.durata(v) for v in valori] + [esecuzione.durata(v) for v in valori]
                              + [scelte.gb(v) for v in (0, 0.3, 0.6, 12, 1234.5)]
                              + [numero(v, c) for v in (0, 15.6, 1234567.891) for c in (0, 1, 2, 3)]
                              + [sicurezza.due_cifre(v) for v in (0, 1, 16.95, 30023, 0.0456)])
    finally:
        for w in finestre:
            w.close()
            w.deleteLater()
        progetto.cartella_progetti, finestra.AVVIO_MIN_MS = cartella_progetti, avvio_min
        pagina_motore.spazio_windows_gb = spazio_windows
        if lingua_scelta:
            from app import lingua
            lingua.imposta("it")
        for k, v in vecchi.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v
        shutil.rmtree(tmp, ignore_errors=True)
    return foto


# ---------------- finto motore, referto, motore ----------------
def finto(cartella, *argomenti, lingua_scelta=None):
    """Eventi del finto motore (velocità 0), senza le parti che cambiano."""
    env = {**os.environ, "TEMP": str(cartella / "temp"), "TMP": str(cartella / "temp")}
    extra = ["--lingua", lingua_scelta] if lingua_scelta else []
    r = subprocess.run([sys.executable, str(FINTO), *argomenti, "--velocita", "0", *extra], capture_output=True,
                       text=True, encoding="utf-8", env=env, timeout=120)
    out = []
    for riga in r.stdout.splitlines():
        e = json.loads(riga)
        e.pop("ora", None)
        e.pop("corsa", None)
        e.pop("log", None)
        if e.get("tipo") == "fine" and isinstance(e.get("risultato"), dict):
            ris = dict(e["risultato"])
            for k in ("cartella", "file", "corsa"):
                ris.pop(k, None)
            if "previsioni" in ris:
                for p in ris["previsioni"]:
                    p.pop("curva", None)
            e["risultato"] = ris
        out.append(pulisci(json.dumps(e, ensure_ascii=False, sort_keys=True), [cartella]))
    return out


EXCEL = r"""
import json, sys
from pathlib import Path
sys.path.insert(0, sys.argv[1])
import referto
from openpyxl import load_workbook
risultati, cfg, uscita = Path(sys.argv[2]), json.loads(sys.argv[3]), Path(sys.argv[4])
referto.excel(referto.raccogli(risultati, cfg), uscita)
wb = load_workbook(uscita)
for ws in wb.worksheets:
    for riga in ws.iter_rows(values_only=True):
        print(json.dumps(f"{ws.title}: " + " | ".join("" if v is None else str(v) for v in riga)))
"""


def python_excel():
    """Un Python con openpyxl (quello dell'app non ce l'ha): questo, o quello di sistema (py), o None."""
    for p in (sys.executable, shutil.which("py"), shutil.which("python3")):
        if p and subprocess.run([p, "-c", "import openpyxl"], capture_output=True).returncode == 0:
            return p
    return None


def testi_excel(risultati, cfg, lingua_scelta=None):
    """Le celle dell'Excel del referto, scritto da referto.excel() con il Python che ha openpyxl."""
    p = python_excel()
    if p is None:
        return ["(nessun Python con openpyxl: Excel non controllato)"]
    env = {**os.environ, **({"PMN_LINGUA": lingua_scelta} if lingua_scelta else {})}
    r = subprocess.run([p, "-c", EXCEL, str(RADICE / "motore"), str(risultati), json.dumps(cfg),
                        str(risultati / "prova_excel.xlsx")], capture_output=True, text=True, encoding="utf-8",
                       env=env, timeout=120)
    if r.returncode != 0:
        raise RuntimeError(r.stderr)
    return [json.loads(x) for x in r.stdout.splitlines()]


def foto_motore(lingua_scelta=None):
    """{caso: [testi]} del finto motore, del referto e dei testi del motore vero."""
    import importlib
    tmp = Path(tempfile.mkdtemp(prefix="pmn_lingua_m_"))
    (tmp / "temp").mkdir()
    foto = {}
    try:
        cfg = tmp / "prog" / "config.yaml"
        cfg.parent.mkdir()
        cfg.write_text("project_name: prog\nmotore:\n  kraken_db: completo\n  virus: true\n", encoding="utf-8")
        cfg_rapido = tmp / "rapido" / "config.yaml"
        cfg_rapido.parent.mkdir()
        cfg_rapido.write_text("project_name: rapido\nmotore:\n  kraken_db: rapido\n", encoding="utf-8")
        L = lingua_scelta

        def f(*a):
            return finto(tmp, *a, lingua_scelta=L)

        foto["m01 info"] = f("info")
        foto["m02 stima"] = f("stima", "--progetto", str(cfg))
        foto["m03 stima tempo lungo"] = f("stima", "--progetto", str(cfg), "--scenario", "TEMPO_LUNGO")
        foto["m04 stima senza progetto"] = f("stima")
        foto["m05 verifica"] = f("verifica")
        foto["m06 verifica rossa"] = f("verifica", "--scenario", "DB_MANCANTE")
        foto["m07 installa"] = f("installa")
        foto["m08 installa annullata"] = f("installa", "--scenario", "annullata")
        foto["m09 annulla"] = f("annulla", "--corsa", "x")
        for sc in ("DB_MANCANTE", "READ_NON_TROVATE", "annullata", "GIA_IN_CORSO", "CAMPIONE_DA_CONTROLLARE",
                   "TEMPO_LUNGO"):
            foto[f"m10 esegui {sc}"] = f("esegui", "--progetto", str(cfg_rapido), "--scenario", sc)
        foto["m11 esegui senza progetto"] = f("esegui")
        foto["m12 referto senza risultati"] = f("referto", "--progetto", str(cfg))
        foto["m13 esegui rapido"] = f("esegui", "--progetto", str(cfg_rapido))
        foto["m14 esegui completo"] = f("esegui", "--progetto", str(cfg))
        richieste = {
            "salsiccia": {"prodotto": {"temperatura": 5, "sale_fase_acquosa": 3.5, "ph": 6.0, "nitrito": 100},
                          "giorni": 21},
            "lattici": {"prodotto": {"temperatura": 5, "sale_fase_acquosa": 3.5, "ph": 6.0}, "giorni": 21,
                        "conte": {"lab_ufc_g": 1000000}},
            "totale": {"prodotto": {"temperatura": 5, "aw": 0.97, "ph": 6.0}, "giorni": 21,
                       "conte": {"totale_ufc_g": 10000000}},
            "salame": {"modalita": "fssp", "prodotto": {"temperatura": 15, "aw": 0.90, "ph": 5.2, "nitrito": 100,
                                                        "co2": 20, "fenoli": 5, "acidi": {"lattico": 7000}},
                       "profilo_temperatura": [[15, 48], [12, 400]], "lag": True},
            "fuori": {"prodotto": {"temperatura": 80, "sale_fase_acquosa": 3.5, "ph": 6.0}, "giorni": 21},
            "testo": {"prodotto": {"temperatura": "cinque"}, "giorni": 21},
            "aw_e_sale": {"prodotto": {"aw": 0.97, "sale_fase_acquosa": 3}, "giorni": 21},
            "campo": {"prodotto": {"temperatura": 5, "pepe": 1}, "giorni": 21},
            "acido": {"prodotto": {"temperatura": 5, "acidi": {"fumarico": 1}}, "giorni": 21},
            "nome": {"prodotto": {"temperatura": 5}, "giorni": 21, "campione": "../x"},
            "giorni": {"prodotto": {"temperatura": 5}, "giorni": 900},
            "profilo": {"prodotto": {"temperatura": 5}, "profilo_specie": {
                "s__Latilactobacillus_sakei": 30, "s__Leuconostoc_carnosum": 5, "s__Listeria_monocytogenes": 0.01},
                "conte": {"totale_ufc_g": 1000000}},
        }
        for nome, richiesta in richieste.items():
            ingresso = tmp / f"{nome}.json"
            ingresso.write_text(json.dumps(richiesta), encoding="utf-8")
            con = ["--progetto", str(cfg)] if nome != "profilo" else []
            foto[f"m20 sicurezza {nome}"] = f("sicurezza", "--ingresso", str(ingresso), *con)
        (tmp / "prog" / "risultati" / "sicurezza" / "SRR4875279.json").unlink()
        ingresso = tmp / "salsiccia.json"
        f("sicurezza", "--ingresso", str(ingresso), "--progetto", str(cfg))
        foto["m21 sicurezza senza file"] = f("sicurezza", "--ingresso", str(tmp / "nessuno.json"))
        # triage nel referto: righe di steps/17
        (tmp / "prog" / "risultati" / "tabelle" / "triage.tsv").write_text(
            "sample\tdecision\thost_pct\test_nonhost_fragments\tq30_pct\tscore\n"
            "SRR4875279\treview\t97\t900000\t60\t0.45\nALTRO\texclude\tx\t\t\t\n", encoding="utf-8")
        foto["m22 referto"] = f("referto", "--progetto", str(cfg))
        referto = tmp / "prog" / "risultati" / "referto"
        foto["m23 referto html"] = [pulisci(r, [tmp]) for r in (referto / "referto.html").read_text(
            encoding="utf-8").splitlines()]
        foto["m24 referto excel"] = [pulisci(r, [tmp]) for r in testi_excel(
            tmp / "prog" / "risultati", {"project_name": "prog", "motore": {"kraken_db": "completo", "virus": True}},
            L)]
        # parametri scelti a mano (non i consigliati) e screening rapido
        par = tmp / "prog" / "risultati" / "tabelle" / "parameters.tsv"
        testo = par.read_text(encoding="utf-8")
        testo = re.sub(r"(?m)^kraken_confidence\t.*$", "kraken_confidence\t0.3", testo)
        testo = re.sub(r"(?m)^metaphlan_mode\t.*$", "metaphlan_mode\tstandard", testo)
        par.write_text(testo, encoding="utf-8")
        f("referto", "--progetto", str(cfg))
        foto["m25 referto parametri scelti"] = [pulisci(r, [tmp]) for r in (referto / "referto.html").read_text(
            encoding="utf-8").splitlines() if "<td" in r or "<tr" in r]

        # motore vero: testi che non passano dal finto
        import pmn_motore
        import referto as ref
        import system_one
        tr = lambda t: t  # noqa: E731  (i testi fissi del motore, come li mostra: tradotti con la lingua scelta)
        if L:
            import lingua as lm
            lm.imposta(L)
            tr = lm._
        try:
            righe = [pmn_motore.messaggio_fine("Messaggio.", "Cosa fare.", p) for p in
                     list(pmn_motore.NOMI_PASSI) + [None, "ignoto"]]
            righe += [f"{p}: {tr(m)}" for p, m in pmn_motore.PASSI]
            righe += [f"{c}: {tr(m)} / {tr(x)}" for _, c, m, x in pmn_motore.ERRORI]
            righe += [tr(c) for v in pmn_motore.CITAZIONI.values() for c in v]
            righe += [tr(pmn_motore.ESEMPIO["frase"])]
            for ram in (4.0, 15.6, 64.0):
                for kc, kr in ((True, True), (False, True), (False, False)):
                    _, cons = system_one.modo(ram, ram > 40, kc, kr)
                    righe += [c["messaggio"] + " / " + c["fonte"] for c in cons]
            righe += [c["messaggio"] for c in system_one.tempo({"kraken": 20000})]
            righe += [c["messaggio"] for c in system_one.spazio(100, 40)]
            righe += [c["messaggio"] for c in system_one.database("rapido", {})]
            righe += [f"{p['frase']} / {p['perche']} / {p['fonte']}" for p in system_one.prodotti()]
            tsv = tmp / "triage.tsv"
            tsv.write_text("sample\tdecision\thost_pct\test_nonhost_fragments\tq30_pct\tscore\n"
                           "A\tpass\t10\t5000000\t90\t0.9\nB\treview\t97\t900000\t60\t0.45\nC\texclude\tx\t\t\t\n",
                           encoding="utf-8")
            esiti, cons = system_one.triage(tsv)
            righe += [json.dumps(e, ensure_ascii=False) for e in esiti] + [c["messaggio"] + " / " + c["fonte"]
                                                                           for c in cons]
            righe += [ref.problema(x, m, e) for x, m, e in (("l'Excel", "openpyxl", ImportError()),
                                                            ("il PDF", "PySide6", PermissionError()),
                                                            ("il PDF", "PySide6", RuntimeError()))]
            righe += [ref.avviso_parziale(["manca il PDF"], True), ref.avviso_parziale([], False)]
            r = subprocess.run([sys.executable, str(RADICE / "motore" / "pmn_motore.py"), "info",
                                *(["--lingua", L] if L else [])],
                               capture_output=True, text=True, encoding="utf-8", timeout=120,
                               env={**os.environ, "PMN_HOME": str(tmp / "senza_installazione")})
            for riga in r.stdout.splitlines():
                e = json.loads(riga)
                righe.append(f"{e['tipo']}: {e.get('messaggio', '')}")
            foto["m30 motore vero"] = [pulisci(x, [tmp]) for x in righe]
        finally:
            if L:
                lm.imposta("it")
            importlib.reload(system_one)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return foto


def foto(lingua_scelta=None):
    return {**foto_app(lingua_scelta), **foto_motore(lingua_scelta)}


# parole italiane che in un testo inglese non ci devono essere (parole intere, maiuscole o minuscole)
ITALIANE = set("""il lo la gli di del della dello delle dei degli che non con sono una nel nella nei sul sulla
alla alle al dal dalla da fra tra anche ma questo questa questi quello quelle ho hai può puoi sei siamo essere
campione campioni referto motore analisi specie soglia prodotto progetto progetti risultati cartella scegli premi
apri controlla manca mancano errore avviso pronto pronta finita fermata corso passo passi tempo spazio stima
microbi batteri funghi archei ospite maiale bovino pecora uomo salume salsiccia previsione modello crescita conta
giorni ore minuti secondi circa sopra sotto dopo prima ancora già più meno tutte tutti altro altri nessuno qui
profilo sicurezza alimentare lettura letture coppie valore fonte parametri strumenti""".split())
# testi che restano uguali in tutte le lingue (nomi propri, valori del file del progetto): non sono frasi italiane
AMMESSE = ("[file] ", "Felino", "fette", "rapido", "completo")
# messaggi che questa fotografia fa arrivare da sé alle pagine (dati di prova, non testi dell'app)
DI_PROVA = ("Manca il database di Kraken.", "Apri la pagina Motore e premi Installa o completa.", "Poca RAM.", "Messaggio.",
            "Cosa fare.", "manca il PDF", "ALTRO")
# negli eventi del motore (JSON) l'utente vede solo questi campi; gli altri sono chiavi del contratto (passo, esito...)
VISIBILI = {"messaggio", "cosa_fare", "frase", "perche", "note", "profilo_usato", "citazione", "citazioni", "scenario",
            "motivi", "nome"}


def da_mostrare(testo):
    """I testi per l'utente dentro una riga della fotografia: la riga stessa, o i campi VISIBILI se è un evento JSON."""
    try:
        d = json.loads(testo)
    except ValueError:
        return [testo]
    out = []

    def giro(x, chiave=None):
        if isinstance(x, dict):
            for k, v in x.items():
                giro(v, k)
        elif isinstance(x, list):
            for v in x:
                giro(v, chiave)
        elif isinstance(x, str) and chiave in VISIBILI:
            out.append(x)
    giro(d)
    return out


def italiane(testo):
    """Le parole italiane di un testo (vuoto se non ce ne sono). Non contano: i nomi con _ o / (progetti di prova,
    file, chiavi), i segnaposto come <CARTELLA>, i nomi propri e i valori di AMMESSE, i messaggi DI_PROVA."""
    if testo.startswith(AMMESSE[0]):
        return []
    for x in DI_PROVA + AMMESSE[1:]:
        testo = testo.replace(x, "")
    testo = re.sub(r"<[A-Z]+>|\S*[_/]\S*|et al\.|Le Marc|Command \w+", " ", testo)   # Command info: nome del comando
    testo = re.sub(r"^(?:ospite|inizio|avanzamento|avviso|errore|fine): ", "", testo)       # chiave del contratto: testo
    parole = re.findall(r"[A-Za-zÀ-ÿ'-]+", testo)
    return [p for p in parole if p.lower().strip("'") in ITALIANE or re.search(r"[àèéìòù]", p.lower())]


if __name__ == "__main__":
    if sys.argv[1:] == ["--scrivi"]:
        FOTO_IT.parent.mkdir(exist_ok=True)
        f = foto()
        FOTO_IT.write_text(json.dumps(f, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"{FOTO_IT}: {len(f)} casi, {sum(len(v) for v in f.values())} testi")
    else:
        print(__doc__)
